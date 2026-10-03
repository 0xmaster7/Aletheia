"""Extract content and images from the Word template."""
import os
from docx import Document
from docx.oxml.ns import qn
import zipfile

DOCX_PATH = "/Users/keshavnanda/Desktop/Aletheia/3 BCSE497J Project I Report - Template.docx"
OUT_DIR = "/Users/keshavnanda/Desktop/Aletheia/report_extract"
IMG_DIR = os.path.join(OUT_DIR, "figures")
os.makedirs(IMG_DIR, exist_ok=True)

# 1. Extract images from docx (it's a zip archive)
with zipfile.ZipFile(DOCX_PATH, 'r') as z:
    img_count = 0
    for name in z.namelist():
        if name.startswith("word/media/"):
            ext = os.path.splitext(name)[1]
            img_count += 1
            out_name = f"fig{img_count}{ext}"
            with z.open(name) as src, open(os.path.join(IMG_DIR, out_name), 'wb') as dst:
                dst.write(src.read())
            print(f"  Extracted: {out_name} (from {name})")
    print(f"\nTotal images extracted: {img_count}\n")

# 2. Extract text content with structure
doc = Document(DOCX_PATH)
print("=" * 80)
print("DOCUMENT STRUCTURE")
print("=" * 80)

for i, para in enumerate(doc.paragraphs):
    style = para.style.name if para.style else "None"
    text = para.text.strip()
    if text:
        # Check for bold/italic
        runs_info = []
        for run in para.runs:
            bold = "B" if run.bold else ""
            italic = "I" if run.italic else ""
            underline = "U" if run.underline else ""
            fmt = f"[{bold}{italic}{underline}]" if any([bold, italic, underline]) else ""
            if run.text.strip():
                runs_info.append(f"{fmt}{run.text.strip()}")
        
        print(f"\n[P{i:03d}] Style='{style}'")
        print(f"  Text: {text[:200]}")
        if runs_info:
            print(f"  Runs: {' | '.join(runs_info[:5])}")

# 3. Extract tables
print(f"\n\n{'=' * 80}")
print("TABLES")
print("=" * 80)
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx + 1} ({len(table.rows)} rows × {len(table.columns)} cols) ---")
    for r_idx, row in enumerate(table.rows):
        cells = [cell.text.strip().replace('\n', ' | ') for cell in row.cells]
        print(f"  Row {r_idx}: {cells}")

# 4. Check sections for page layout
print(f"\n\n{'=' * 80}")
print("PAGE LAYOUT")
print("=" * 80)
for s_idx, section in enumerate(doc.sections):
    print(f"Section {s_idx}:")
    print(f"  Page width:  {section.page_width.inches:.2f} in")
    print(f"  Page height: {section.page_height.inches:.2f} in")
    print(f"  Left margin: {section.left_margin.inches:.2f} in")
    print(f"  Right margin: {section.right_margin.inches:.2f} in")
    print(f"  Top margin:  {section.top_margin.inches:.2f} in")
    print(f"  Bottom margin: {section.bottom_margin.inches:.2f} in")
